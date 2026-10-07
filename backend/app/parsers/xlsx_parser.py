import io
import zipfile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.domain.structure import ParseContext, WorkbookStructure
from app.errors import IngestionError
from app.messages import text
from app.parsers.common import build_sheet, workbook_from_sheets


def assert_zip_container_safe(
    content: bytes,
    max_uncompressed_bytes: int,
    max_entries: int = 5000,
) -> None:
    if not content.startswith(b"PK"):
        raise IngestionError("unreadable_workbook", text("unreadable_workbook"))
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            infos = archive.infolist()
            if len(infos) > max_entries:
                raise IngestionError("uncompressed_too_large", text("uncompressed_too_large"))
            total = 0
            for info in infos:
                total += info.file_size
                if total > max_uncompressed_bytes:
                    raise IngestionError("uncompressed_too_large", text("uncompressed_too_large"))
    except IngestionError:
        raise
    except (zipfile.BadZipFile, OSError) as exc:
        raise IngestionError("unreadable_workbook", text("unreadable_workbook")) from exc


def parse_xlsx(content: bytes, context: ParseContext) -> WorkbookStructure:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    except (InvalidFileException, zipfile.BadZipFile, OSError, ValueError) as exc:
        raise IngestionError("unreadable_workbook", text("unreadable_workbook")) from exc
    try:
        try:
            names = [str(name) for name in workbook.sheetnames]
            if len(names) > context.limits.max_sheets:
                raise IngestionError("too_many_sheets", text("too_many_sheets"))
            sheets = [
                build_sheet(name, workbook[name].iter_rows(values_only=True), context.limits)
                for name in names
            ]
            return workbook_from_sheets("xlsx", sheets)
        except IngestionError:
            raise
        except (InvalidFileException, zipfile.BadZipFile, OSError, ValueError, KeyError) as exc:
            raise IngestionError("unreadable_workbook", text("unreadable_workbook")) from exc
    finally:
        workbook.close()
