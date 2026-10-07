import io
from pathlib import Path

import pytest
from app.config import Settings
from app.factory import create_app
from fastapi.testclient import TestClient
from openpyxl import Workbook


def make_settings(tmp_path: Path, **overrides: object) -> Settings:
    data_dir = tmp_path / "data"
    values: dict[str, object] = {
        "data_dir": data_dir,
        "database_url": f"sqlite:///{(tmp_path / 'test.sqlite3').as_posix()}",
        "max_upload_bytes": 1024 * 1024,
        "max_files_per_request": 10,
        "max_sheets": 20,
        "max_rows_per_sheet": 1000,
        "max_columns": 64,
        "max_uncompressed_bytes": 5 * 1024 * 1024,
        "cors_origins": ("http://localhost:3000",),
    }
    values.update(overrides)
    return Settings(**values)  # type: ignore[arg-type]


def xlsx_bytes(sheets: dict[str, list[list[object]]]) -> bytes:
    workbook = Workbook()
    default_sheet = workbook.active
    if default_sheet is None:
        raise RuntimeError("workbook has no active sheet")
    first = True
    for name, rows in sheets.items():
        worksheet = default_sheet if first else workbook.create_sheet(name)
        first = False
        worksheet.title = name
        for row in rows:
            worksheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def xls_bytes() -> bytes:
    import xlwt

    book = xlwt.Workbook(encoding="utf-8")
    sheet = book.add_sheet("فروش")
    sheet.write(0, 0, "نام")
    sheet.write(0, 1, "مبلغ")
    sheet.write(1, 0, "علی")
    sheet.write(1, 1, 10)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def upload(client: TestClient, filename: str, content: bytes):
    return client.post(
        "/api/uploads",
        files=[("files", (filename, content, "application/octet-stream"))],
    )


@pytest.fixture
def settings(tmp_path: Path) -> Settings:
    return make_settings(tmp_path)


@pytest.fixture
def client(settings: Settings):
    application = create_app(settings)
    with TestClient(application) as test_client:
        yield test_client
